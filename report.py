"""Write the Excel tracker an analyst would open."""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILL = {"green": "C6EFCE", "amber": "FFEB9C", "red": "FFC7CE"}
SCENARIO_TEXT = {
    "clean": "Complete, consistent file",
    "dob_mismatch": "DOB on subscription differs from ID",
    "name_mismatch": "Name on subscription differs from ID",
    "address_mismatch": "Address on subscription differs from ID",
    "missing_dob": "DOB left blank on subscription",
}


def scenario_text(truth, expired):
    text = SCENARIO_TEXT[truth["scenario"]]
    return text + (" + ID expired" if expired else "")


def style(sheet, widths, light_col=None):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if light_col:
            cell = row[light_col - 1]
            if cell.value and str(cell.value).lower() in FILL:
                cell.fill = PatternFill("solid", fgColor=FILL[str(cell.value).lower()])
    sheet.freeze_panes = "A2"


def method_table(method_rows):
    """Accuracy per document, method, condition and field."""
    df = pd.DataFrame(method_rows)
    table = df.groupby(["document", "method", "condition", "field"])["correct"].agg(["sum", "count"]).reset_index()
    table.columns = ["document", "method", "condition", "field", "correct", "total"]
    table["accuracy"] = table["correct"] / table["total"]
    return table


def write_tracker(results, metrics, method_rows, path, expired_ids):
    wb = Workbook()

    ws = wb.active
    ws.title = "Investor tracker"
    ws.append(["Investor", "Status", "Documents found", "Flags for analyst", "Answer key (synthetic scenario)"])
    for r in results:
        docs = ", ".join(f"{f['name']}: {f['predicted']}" for f in r["files"])
        flags = [f"[{x['light'].upper()}] {x['document']} / {x['field']}: {x['note']}"
                 for x in r["records"] if x["light"] != "green"]
        ws.append([r["truth"]["investor_id"], r["status"].upper(), docs, "\n".join(flags),
                   scenario_text(r["truth"], r["truth"]["investor_id"] in expired_ids)])
    style(ws, [11, 9, 34, 80, 34], light_col=2)

    ws = wb.create_sheet("Field detail")
    ws.append(["Investor", "Document", "Field", "Value", "Light", "Sources", "Note", "Correct?"])
    for r in results:
        for x in r["records"]:
            ws.append([x["investor_id"], x["document"], x["field"], x["value"], x["light"].upper(),
                       ", ".join(x["sources"]), x["note"],
                       {True: "yes", False: "NO", None: "-"}[x.get("correct")]])
    style(ws, [11, 13, 18, 34, 8, 18, 50, 9], light_col=5)

    ws = wb.create_sheet("Method comparison")
    ws.append(["Document", "Method", "Condition", "Field", "Correct", "Total", "Accuracy"])
    for row in method_table(method_rows).itertuples(index=False):
        ws.append([row.document, row.method, row.condition, row.field, int(row.correct), int(row.total),
                   f"{row.accuracy:.0%}"])
    style(ws, [13, 14, 12, 18, 9, 7, 10])

    ws = wb.create_sheet("Evaluation")
    ws.append(["Metric", "Result"])
    for k, v in metrics.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["IDs: DocXPand-25k (synthetic, CC BY-NC-SA). Subscription forms: blank SEC EDGAR exhibits "
               "filled with synthetic data."])
    style(ws, [52, 14])

    wb.save(path)
